"""
excel_import.py

Reads the "Dimensions and Placement" sheet from the spacecraft systems workbook
and exposes each row as a ComponentSpec object.

The workbook is treated as the single source of truth.  Dimension cells in that
sheet may themselves be formulas linked to the subsystem specification sheets;
therefore the workbook is loaded with data_only=True so Python reads the cached
calculated values saved by Excel.

If you edit formulas in the workbook, open/save the workbook in Excel first so
the cached values are refreshed before rerunning the CAD script.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Iterable, Optional, Tuple

from openpyxl import load_workbook


SHEET_NAME = "Dimensions and Placement"


@dataclass(frozen=True)
class ComponentSpec:
    subsystem: str
    name: str
    model: str
    dimensions_mm: Tuple[Optional[float], Optional[float], Optional[float]]
    position_mm: Tuple[Optional[float], Optional[float], Optional[float]]
    orientation: Tuple[Optional[float], Optional[float], Optional[float]]
    design_status: str = ""
    cad_notes: str = ""
    source: str = ""

    @property
    def length(self) -> Optional[float]:
        return self.dimensions_mm[0]

    @property
    def width(self) -> Optional[float]:
        return self.dimensions_mm[1]

    @property
    def height(self) -> Optional[float]:
        return self.dimensions_mm[2]

    @property
    def x(self) -> Optional[float]:
        return self.position_mm[0]

    @property
    def y(self) -> Optional[float]:
        return self.position_mm[1]

    @property
    def z(self) -> Optional[float]:
        return self.position_mm[2]


def _number_or_none(value) -> Optional[float]:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return float(value)
    if isinstance(value, (int, float)):
        return float(value)
    return None


def load_components(
    workbook_path: str | Path,
    sheet_name: str = SHEET_NAME,
) -> Dict[str, ComponentSpec]:
    """
    Load all populated component rows from the placement sheet.

    Returns
    -------
    dict
        Dictionary keyed by the exact component name in column B.
    """
    workbook_path = Path(workbook_path)
    if not workbook_path.exists():
        raise FileNotFoundError(f"Workbook not found: {workbook_path}")

    wb = load_workbook(workbook_path, data_only=True, read_only=False)

    if sheet_name not in wb.sheetnames:
        raise KeyError(
            f'Sheet "{sheet_name}" not found. Available sheets: {wb.sheetnames}'
        )

    ws = wb[sheet_name]
    components: Dict[str, ComponentSpec] = {}

    # Row 2 is the column header row in the current master workbook.
    for row in range(3, ws.max_row + 1):
        subsystem = ws.cell(row, 1).value
        name = ws.cell(row, 2).value

        if not subsystem or not name:
            continue

        spec = ComponentSpec(
            subsystem=str(subsystem).strip(),
            name=str(name).strip(),
            model=str(ws.cell(row, 3).value or "").strip(),
            dimensions_mm=(
                _number_or_none(ws.cell(row, 4).value),
                _number_or_none(ws.cell(row, 5).value),
                _number_or_none(ws.cell(row, 6).value),
            ),
            position_mm=(
                _number_or_none(ws.cell(row, 7).value),
                _number_or_none(ws.cell(row, 8).value),
                _number_or_none(ws.cell(row, 9).value),
            ),
            orientation=(
                _number_or_none(ws.cell(row, 10).value),
                _number_or_none(ws.cell(row, 11).value),
                _number_or_none(ws.cell(row, 12).value),
            ),
            design_status=str(ws.cell(row, 13).value or "").strip(),
            cad_notes=str(ws.cell(row, 14).value or "").strip(),
            source=str(ws.cell(row, 15).value or "").strip(),
        )

        if spec.name in components:
            raise ValueError(f'Duplicate component name "{spec.name}" in {sheet_name}')

        components[spec.name] = spec

    return components


def require_component(
    components: Dict[str, ComponentSpec],
    name: str,
) -> ComponentSpec:
    """Return a named component or raise a readable error."""
    try:
        return components[name]
    except KeyError as exc:
        available = "\n  - ".join(sorted(components))
        raise KeyError(
            f'Component "{name}" was not found.\nAvailable components:\n  - {available}'
        ) from exc


def require_dimensions(spec: ComponentSpec) -> Tuple[float, float, float]:
    """Return L/W/H and fail early if a required CAD dimension is missing."""
    if any(v is None for v in spec.dimensions_mm):
        raise ValueError(
            f'{spec.name}: missing dimensions {spec.dimensions_mm}. '
            "Check the workbook and ensure Excel has recalculated/saved it."
        )
    return tuple(float(v) for v in spec.dimensions_mm)  # type: ignore[arg-type]


def require_position(spec: ComponentSpec) -> Tuple[float, float, float]:
    """Return X/Y/Z and fail early if a required CAD position is missing."""
    if any(v is None for v in spec.position_mm):
        raise ValueError(
            f'{spec.name}: missing position {spec.position_mm}. '
            "Populate the placement sheet before generating this component."
        )
    return tuple(float(v) for v in spec.position_mm)  # type: ignore[arg-type]


if __name__ == "__main__":
    # Lightweight import check.
    default_book = Path(__file__).with_name(
        "budget_closing_consolidated_master_with_placement_mass.xlsx"
    )
    components = load_components(default_book)

    for name in [
        "Microsatellite bus",
        "Telescope Assembly",
        "Baffle",
        "Tank",
        "Battery",
    ]:
        c = require_component(components, name)
        print(
            f"{c.name:24s} "
            f"dims={c.dimensions_mm} mm  "
            f"pos={c.position_mm} mm  "
            f"ori={c.orientation}"
        )
